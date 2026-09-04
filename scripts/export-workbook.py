"""Read-only, reconciled export of the specific RIVINTER journal layout.
Usage: python scripts/export-workbook.py source.xlsx private-output.json
Never refresh external links or overwrite the original workbook.
"""
import sys, json, hashlib, datetime
from pathlib import Path
import openpyxl

source, target = map(Path, sys.argv[1:3])
if source.resolve() == target.resolve():
    raise ValueError("The source must never be overwritten")
raw = openpyxl.load_workbook(source, data_only=False)
cached = openpyxl.load_workbook(source, data_only=True)
s, formulas = cached["RIVINTER"], raw["RIVINTER"]
types = ["B65", "B33N", "B33V", "B30CL", "ALE50", "BAC"]
prices = dict(zip(types, [16500, 16500, 16500, 16500, 24500, 4500]))
cols = dict(zip(types, [4, 7, 10, 13, 16, 19]))
def serial(v):
    if isinstance(v, openpyxl.worksheet.formula.ArrayFormula):
        return {"type":"array_formula","text":v.text,"ref":v.ref}
    return v.isoformat() if isinstance(v, (datetime.date, datetime.datetime, datetime.time)) else v
def number(v):
    if v is None:
        return 0
    if not isinstance(v, (int, float)) or int(v) != v:
        raise ValueError(f"Non-integer accounting value: {v!r}")
    return int(v)
opening = {t:number(s.cell(6,c+2).value) for t,c in cols.items()}
running = opening.copy()
assert sum(opening[t]*prices[t] for t in types) == number(s["V6"].value)
kinds = {"BON DE CONSIGNATION":"delivery", "BON DE DECONSIGNATION":"return",
         "CONSIGNATION PAYANTE":"paid", "CONSIGNATION BOUTEILLES":"bottles",
         "CASSE ET MANQUANT BOUTEILLES":"breakage", "ANNUL. BON DE DECONSIGN.":"legacy"}
operations = []
for row in range(7,180):
    ref = s.cell(row,3).value
    if ref is None:
        raise ValueError(f"Missing reference in expected journal row {row}")
    label = str(s.cell(row,2).value).strip()
    if label not in kinds:
        raise ValueError(f"Unrecognized source label at row {row}: {label!r}")
    date = s.cell(row,1).value
    date_iso = date.date().isoformat() if isinstance(date,datetime.datetime) else date.isoformat() if isinstance(date,datetime.date) else None
    debit = {t:number(s.cell(row,c).value) for t,c in cols.items()}
    credit = {t:number(s.cell(row,c+1).value) for t,c in cols.items()}
    lines = [{"bremer_id":t,"delta":debit[t]-credit[t],"unit_price":prices[t]} for t in types]
    for t,c in cols.items():
        running[t] += debit[t]-credit[t]
        if running[t] != number(s.cell(row,c+2).value):
            raise ValueError(f"Running balance mismatch at row {row}, {t}")
    value = sum(running[t]*prices[t] for t in types)
    if value != number(s.cell(row,22).value):
        raise ValueError(f"Value mismatch at row {row}: {value}")
    operations.append({"source_key":f"RIVINTER:{row}","date":date_iso,"kind":kinds[label],"ref":str(ref),
      "lines":lines,"metadata":{"row":row,"source_kind":label,"original_date":serial(date),
      "original_debits":debit,"original_credits":credit,"original_balances":running.copy(),"original_value":value}})
archive = {"format":"rivinter-workbook-v1","file_hash":hashlib.sha256(source.read_bytes()).hexdigest(),
 "source_filename":source.name,"sheet":"RIVINTER","range":"A6:W179","external_links_refreshed":False,
 "opening":{"quantities":opening,"source_metadata":{"raw_date":serial(s["A6"].value),"raw_label":s["B6"].value,
 "raw_ref":s["C6"].value,"warning":"Date du report contradictoire : A6, libellé et mois différents. Confirmer la date sans modifier les soldes."}},
 "operations":operations,"expected_closing":{"quantities":running,"value":value,"count":len(operations)},
 "original_cells":[{"cell":cell.coordinate,"value":serial(s[cell.coordinate].value),"original":serial(cell.value)}
    for row in formulas.iter_rows() for cell in row if cell.value is not None],
 "notes":["Les cellules hors A6:W179 sont archivées mais ne deviennent pas des mouvements.",
 "Les dates #N/A sont conservées et importées comme dates non confirmées.",
 "Les références BV répétées et les écritures à zéro sont conservées par numéro de ligne."]}
target.parent.mkdir(parents=True, exist_ok=True)
if target.exists():
    raise ValueError("Refusing to overwrite an existing export")
target.write_text(json.dumps(archive,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps({"output":str(target),"count":len(operations),"undated":sum(o["date"] is None for o in operations),
 "closing":running,"value":value,"reconciled_rows":len(operations)},ensure_ascii=False))
